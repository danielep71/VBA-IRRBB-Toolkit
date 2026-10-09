#!/usr/bin/env python3
"""Build the development workbook (regression harness and its cases).

Usage: python tools/devbuild/build_dev_workbook.py [output.xlsm]
Default output: build/IRRBB_Dev.xlsm (ignored by Git; never commit it).

The workbook is assembled from one commit: every .bas in src/core, src/modules,
src/workbook, tests/modules and examples/modules is embedded unchanged (CRLF),
with document modules for ThisWorkbook and each sheet, so Excel compiles the
project from source on open. It is the owner-chosen import route of issue #6
and the vehicle for the harness evidence of issue #8; it is not the workbook
template of issue #5. Requires openpyxl.
"""
from __future__ import annotations

import io
import re
import subprocess
import sys
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).resolve().parent))
from vba_project import Module, build_vba_project, document_module  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
NAVY, GREY = "1F3A5F", "F2F4F7"
HDR = (Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY))
THIN = Side(style="thin", color="C8CED6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SHEETS = [("Readme", "shReadme"), ("Checks", "shChecks")]
COMPONENT_DIRS = ["src/core", "src/modules", "src/workbook", "tests/modules", "examples/modules"]
BUTTONS = {
    "Checks": [("Run tests", "RunTests", 1, 3, "2E7D7A")],
}


def git(*args: str) -> str:
    return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()


def header(ws, row: int, col: int, labels: list[str]) -> None:
    for i, text in enumerate(labels):
        c = ws.cell(row, col + i, text)
        c.font, c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def components() -> list[Path]:
    out = []
    for d in COMPONENT_DIRS:
        out += sorted((ROOT / d).glob("*.bas")) if (ROOT / d).is_dir() else []
    return out


def build(target: Path) -> None:
    sha = git("rev-parse", "HEAD")
    # Any tracked change, or an untracked file where components are collected (it would
    # be embedded without being part of the commit), makes the build not evidence-grade.
    dirty = bool(git("status", "--porcelain", "--untracked-files=no")
                 or git("status", "--porcelain", "--untracked-files=all", "--", *COMPONENT_DIRS))
    built = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    comps = components()

    wb = Workbook()
    wb.remove(wb.active)
    ws_by = {}
    for name, code in SHEETS:
        ws = wb.create_sheet(name)
        ws.sheet_properties.codeName = code
        ws.sheet_view.showGridLines = False
        ws_by[name] = ws
    wb.code_name = "ThisWorkbook"

    # Readme ------------------------------------------------------------------
    ws = ws_by["Readme"]
    ws["A1"] = "IRRBB Toolkit - development workbook (regression harness)"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = "Built from source for compile and regression evidence. Not the template of issue #5; never commit it."
    ws["A2"].font = Font(italic=True, color="5B6573")
    info = [("Source commit", sha + (" (dirty working tree: NOT evidence-grade)" if dirty else "")),
            ("Built (UTC)", built), ("Builder", "tools/devbuild/build_dev_workbook.py"),
            ("Components", ", ".join(p.stem for p in comps)),
            ("References", "Default four only (VBA, Excel, OLE Automation, Office)")]
    for i, (k, v) in enumerate(info):
        ws.cell(4 + i, 1, k).font = Font(bold=True)
        ws.cell(4 + i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    steps = [
        "Evidence steps (docs/EXCEL_EVIDENCE.md, INSTALLATION.md):",
        "1. Unblock the file if downloaded (Properties > Unblock), open it and enable macros.",
        "2. Alt+F11 > Tools > References: only the four defaults must be ticked.",
        "3. Debug > Compile VBAProject: must complete without error.",
        "4. Sheet Checks > Run tests: every built case must be PASS; NOT RUN cases are listed, never counted as passed.",
        "5. Run tests a second time: the status must show RUNNING while it runs and the tables must be rewritten, not appended.",
        "6. Record commit, Excel version/build/bitness, OS, compile result and harness counts (docs/EXCEL_EVIDENCE.md).",
    ]
    for i, t in enumerate(steps):
        ws.cell(11 + i, 1, t).font = Font(bold=(i == 0), color=NAVY if i == 0 else "000000")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 110

    # Checks -------------------------------------------------------------------
    ws = ws_by["Checks"]
    ws["A1"] = "Regression harness (TEST_Harness.RunTests)"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    labels = ["Status", "Started", "Finished", "Source commit", "Host", "Cases", "Assertions",
              "Primary error", "Cleanup error"]
    for i, lab in enumerate(labels):
        ws.cell(3 + i, 1, lab).font = Font(bold=True)
        c = ws.cell(3 + i, 2)
        c.fill = PatternFill("solid", fgColor=GREY)
        c.border = BOX
    ws["B3"] = "NOT RUN"
    for cell in ("B4", "B5"):
        ws[cell].number_format = "yyyy-mm-dd hh:mm:ss"
    ws.conditional_formatting.add("B3", CellIsRule(operator="equal",
                                                   formula=['"COMPLETE - ALL EXECUTED CASES PASSED"'],
                                                   fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add("B3", FormulaRule(formula=['OR(LEFT(B3,6)="FAILED",B3="COMPLETE - FAILURES")'],
                                                    fill=PatternFill("solid", fgColor="FFC7CE")))
    header(ws, 13, 1, ["Case", "Outcome", "Assertions", "Passed", "Failed", "Note"])
    header(ws, 13, 8, ["Case", "Assertion", "Expected", "Actual", "AbsDiff", "Tolerance", "Outcome"])
    ws.conditional_formatting.add("B14:B200", CellIsRule(operator="equal", formula=['"PASS"'],
                                                         fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add("B14:B200", CellIsRule(operator="equal", formula=['"FAIL"'],
                                                         fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add("B14:B200", CellIsRule(operator="equal", formula=['"ERROR"'],
                                                         fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add("N14:N20000", CellIsRule(operator="equal", formula=['"FAIL"'],
                                                           fill=PatternFill("solid", fgColor="FFC7CE")))
    for col, w in zip("ABCDEFGHIJKLMN", (16, 14, 11, 9, 9, 60, 3, 14, 34, 18, 18, 10, 14, 9)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["B"].width = 40
    ws.freeze_panes = "A14"

    buf = io.BytesIO()
    wb.save(buf)

    modules = [document_module("ThisWorkbook", True)] + [document_module(code, False) for _, code in SHEETS]
    for path in comps:
        text = path.read_bytes().decode("cp1252").replace("\r\n", "\n").replace("\n", "\r\n")
        modules.append(Module(path.stem, text if text.endswith("\r\n") else text + "\r\n"))
    project_id = "{" + str(uuid.uuid5(uuid.NAMESPACE_URL, "irrbb-dev:" + sha)).upper() + "}"
    inject(buf.getvalue(), build_vba_project(modules, project_id), target, [n for n, _ in SHEETS])


def button_anchor(idx: int, label: str, macro: str, col: int, row: int, color: str) -> str:
    return (f'<xdr:twoCellAnchor editAs="absolute">'
            f'<xdr:from><xdr:col>{col + 1}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
            f'<xdr:to><xdr:col>{col + 1}</xdr:col><xdr:colOff>1600000</xdr:colOff><xdr:row>{row + 2}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
            f'<xdr:sp macro="[0]!{macro}" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{1000 + idx}" name="btn{macro}"/><xdr:cNvSpPr/></xdr:nvSpPr>'
            f'<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm><a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:ln><a:noFill/></a:ln></xdr:spPr>'
            f'<xdr:txBody><a:bodyPr vertOverflow="clip" rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/>'
            f'<a:r><a:rPr lang="en-US" sz="1100" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr>'
            f'<a:t>{label}</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>')


def inject(xlsx: bytes, vba_bin: bytes, out: Path, sheet_names: list[str]) -> None:
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    files = {n: src.read(n) for n in src.namelist()}
    ct = files["[Content_Types].xml"].decode()
    ct = ct.replace("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
    extra = '<Override PartName="/xl/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/>'
    rels = files["xl/_rels/workbook.xml.rels"].decode().replace(
        "</Relationships>", '<Relationship Id="rIdVBA" Type="http://schemas.microsoft.com/office/2006/relationships/'
        'vbaProject" Target="vbaProject.bin"/></Relationships>')
    files["xl/_rels/workbook.xml.rels"] = rels.encode()
    files["xl/vbaProject.bin"] = vba_bin

    idx = 0
    for sheet_no, name in enumerate(sheet_names, 1):
        if name not in BUTTONS:
            continue
        anchors = ""
        for label, macro, col, row, color in BUTTONS[name]:
            idx += 1
            anchors += button_anchor(idx, label, macro, col, row, color)
        sheet_path = f"xl/worksheets/sheet{sheet_no}.xml"
        rel_path = f"xl/worksheets/_rels/sheet{sheet_no}.xml.rels"
        sheet = files[sheet_path].decode()
        rel = files.get(rel_path, b"").decode()
        m = re.search(r'<Relationship[^>]*Type="[^"]*/drawing"[^>]*Target="\.\./drawings/([^"]+)"', rel)
        if m:  # append to the drawing openpyxl already wrote (charts)
            dpath = "xl/drawings/" + m.group(1)
            files[dpath] = files[dpath].decode().replace("</xdr:wsDr>", anchors + "</xdr:wsDr>").encode()
        else:
            dname = f"drawingButtons{sheet_no}.xml"
            files["xl/drawings/" + dname] = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><xdr:wsDr xmlns:xdr="http://schemas.'
                'openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.'
                'org/drawingml/2006/main">' + anchors + "</xdr:wsDr>").encode()
            extra += (f'<Override PartName="/xl/drawings/{dname}" '
                      'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>')
            if not rel:
                rel = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://'
                       'schemas.openxmlformats.org/package/2006/relationships"></Relationships>')
            rel = rel.replace("</Relationships>", '<Relationship Id="rIdBtn" Type="http://schemas.openxmlformats.'
                              f'org/officeDocument/2006/relationships/drawing" Target="../drawings/{dname}"/>'
                              "</Relationships>")
            files[rel_path] = rel.encode()
            if 'xmlns:r=' not in sheet.split(">", 2)[1]:
                sheet = sheet.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/'
                                      'officeDocument/2006/relationships" ', 1)
            sheet, count = re.subn(r"(<pageMargins[^>]*/>)", r'\1<drawing r:id="rIdBtn"/>', sheet, count=1)
            if not count:
                raise RuntimeError(f"cannot place the drawing in {sheet_path}")
            files[sheet_path] = sheet.encode()
    files["[Content_Types].xml"] = ct.replace("</Types>", extra + "</Types>").encode()
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", files.pop("[Content_Types].xml"))
        for name, data in files.items():
            z.writestr(name, data)


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "build" / "IRRBB_Dev.xlsm"
    build(target)
    print("built", target)
